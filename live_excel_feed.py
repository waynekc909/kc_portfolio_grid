import json
import re
import os
import shutil
import tempfile

HTML_FILE = "index.html"
EXCEL_FILE = "KAHF_Consolidated_Fund_Model  Updated.xlsx"

KEYWORD_MAP = {
    "Kronlein": "dumatau_krohnlein",
    "Dunescape": "dunescape",
    "1470": "erf1470",
    "PAM Tree": "pamtree",
    "Turnstone": "turnstone",
    "Mondesa": "mondesa",
    "The Ridge": "rockycrest",
    "C-Breeze": "cbreeze",
    "Tama": "tama",
    "Nkurenkuru Ext 2": "nkurenkuru",
    "Nkurenkuru Ph 1": "lih",
    "Ext 10": "calgrokuumba_otjomuise10",
    "Heaven's Ark": "otjomuise"
}

def get_web_id(excel_name):
    name_lower = str(excel_name).lower()
    for keyword, web_id in KEYWORD_MAP.items():
        if keyword.lower() in name_lower:
            return web_id
    return None

def get_existing_assets(html_path):
    with open(html_path, "r", encoding="utf-8") as f:
        html = f.read()
    pattern = r'(<script id="assets-data" type="application/json">)(.*?)(</script>)'
    match = re.search(pattern, html, re.DOTALL)
    if match:
        return json.loads(match.group(2)), html, match
    return [], html, None

def run_pipeline():
    print("1. Loading base web framework...")
    assets, html, match = get_existing_assets(HTML_FILE)
    if not match:
        print("Error: Could not find assets data in HTML.")
        return
    asset_dict = {a['id']: a for a in assets}

    print("2. Connecting to live Excel model...")
    temp_excel = os.path.join(tempfile.gettempdir(), "temp_fund_model.xlsx")
    shutil.copy2(EXCEL_FILE, temp_excel)

    try:
        import openpyxl
        wb = openpyxl.load_workbook(temp_excel, data_only=True)
        
        # --- TAB 1: FINANCIAL METRICS ---
        ws_perf = wb['2. Project Performance']
        
        # Exact Column Positions from Diagnostic (Row 4)
        header_row = 4
        col_map = {
            'name': 1,      # Col A: Project
            'status': 2,    # Col B: Status
            'dealSize': 6,  # Col F: Approved deal size / facility
            'valuation': 10,# Col J: Carrying value / NAV
            'moic': 12,     # Col L: MOIC / TVPI
            'irr': 13       # Col M: Updated IRR
        }

        # Dynamic scanner override
        for r in range(1, 10):
            row_values = [str(ws_perf.cell(r, c).value or '').strip().lower() for c in range(1, 20)]
            if any(v == 'project' for v in row_values):
                header_row = r
                for c in range(1, 20):
                    val = str(ws_perf.cell(r, c).value or '').strip().lower()
                    if val == 'project': col_map['name'] = c
                    elif val == 'status': col_map['status'] = c
                    elif 'approved deal size' in val: col_map['dealSize'] = c
                    elif 'carrying value' in val or 'nav' in val: col_map['valuation'] = c
                    elif 'moic' in val: col_map['moic'] = c
                    elif 'updated irr' in val: col_map['irr'] = c
                break

        print(f"Extraction Mapping -> Header Row: {header_row}, Columns: {col_map}")

        for r in range(header_row + 1, ws_perf.max_row + 1):
            excel_name = ws_perf.cell(r, col_map['name']).value
            if not excel_name or str(excel_name).strip().startswith("PORTFOLIO"):
                continue

            web_id = get_web_id(excel_name)
            
            if web_id and web_id in asset_dict:
                asset = asset_dict[web_id]
                
                # Status
                status_raw = str(ws_perf.cell(r, col_map['status']).value).strip().lower()
                if status_raw in ['active', 'completed', 'pipeline']:
                    asset['status'] = 'exited' if status_raw == 'completed' else status_raw
                
                # Deal Size
                deal_size = ws_perf.cell(r, col_map['dealSize']).value
                if isinstance(deal_size, (int, float)):
                    asset['dealSize'] = int(deal_size)
                    
                # Valuation / NAV
                nav = ws_perf.cell(r, col_map['valuation']).value
                if isinstance(nav, (int, float)):
                    asset['valuation'] = int(nav)
                    
                # MOIC
                moic = ws_perf.cell(r, col_map['moic']).value
                if isinstance(moic, (int, float)):
                    asset['moicNumeric'] = round(float(moic), 2)
                elif str(moic).strip().lower() == 'n/a':
                    asset['moicNumeric'] = None 
                    
                # IRR
                irr = ws_perf.cell(r, col_map['irr']).value
                if isinstance(irr, (int, float)):
                    val = float(irr)
                    asset['irr'] = round(val * 100, 2) if val < 2.0 else round(val, 2)
                    print(f"Synced {web_id}: IRR = {asset['irr']}%, MOIC = {asset.get('moicNumeric')}")

        # --- TAB 2: OPERATIONAL/DELIVERY METRICS ---
        ws_deliv = wb['3. Delivery Outputs']
        for r in range(5, ws_deliv.max_row + 1):
            excel_name = ws_deliv.cell(r, 1).value
            web_id = get_web_id(excel_name)
            
            if web_id and web_id in asset_dict:
                asset = asset_dict[web_id]
                
                units = ws_deliv.cell(r, 2).value
                if isinstance(units, (int, float)):
                    asset['units'] = int(units)
                    asset['unitsDisplay'] = f"{int(units)} units"
                
                pct = ws_deliv.cell(r, 6).value
                if isinstance(pct, (int, float)):
                    val_pct = float(pct)
                    clean_pct = round(val_pct * 100, 1) if val_pct <= 1.0 else round(val_pct, 1)
                    asset['pct'] = clean_pct
                    asset['statusBadge'] = f"{clean_pct}% COMPLETE" if asset['status'] == 'active' else asset['statusBadge']

        wb.close()

    except Exception as e:
        print(f"Failed to process Excel model: {e}")
        return
    finally:
        if os.path.exists(temp_excel):
            os.remove(temp_excel)

    print("3. Compiling final JSON and injecting to live framework...")
    new_json_str = json.dumps(list(asset_dict.values()), indent=2, ensure_ascii=False)
    new_html = html[:match.start(2)] + "\n" + new_json_str + "\n" + html[match.end(2):]

    with open(HTML_FILE, "w", encoding="utf-8") as f:
        f.write(new_html)
        
    print("✓ Success! Fund model perfectly synced.")

if __name__ == "__main__":
    run_pipeline()